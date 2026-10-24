"""Manuals, SOPs, drawings and alarm lists (SPEC.md 11.4).

Text comes from the PDF text layer only (no OCR). Scanned drawings and images are stored and listed
but marked "no_text". Alarm lists (Excel or CSV) become rows in `alarm_codes`.
"""
import csv
import io
import logging
import re
import uuid
from pathlib import Path

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from . import audit, jobs, machines_service
from .config import get_config
from .models import DOCUMENT_KINDS, AlarmCode, Document, DocumentChunk, Machine, MachineType, User, utcnow

log = logging.getLogger("thunai.documents")

MIMES = {
    "application/pdf": ".pdf",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "text/csv": ".csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/webp": ".webp",
}
CHUNK_CHARS = 1200
MIN_TEXT = 40  # less text than this in a PDF means a scan without a text layer
ALARM_COLUMNS = {
    "code": ("alarm_code", "code", "alarm", "alarm no", "error code"),
    "message": ("message", "description", "alarm message", "text"),
    "cause": ("cause", "possible cause", "reason"),
    "remedy": ("remedy", "action", "corrective action", "solution"),
    "machine_no": ("machine_no", "machine no", "machine"),
    "machine_type": ("machine_type", "machine type", "type"),
}


def _dir() -> Path:
    path = get_config().media_dir / "documents"
    path.mkdir(parents=True, exist_ok=True)
    return path


def path_of(doc: Document) -> Path:
    return get_config().media_dir / doc.file_path


def _scope(db: Session, machine_no: str, machine_type: str) -> tuple[int | None, int | None]:
    if machine_no:
        return machines_service.get_machine(db, machine_no).id, None
    if machine_type:
        mt = db.scalar(select(MachineType).where(func.lower(MachineType.name) == machine_type.strip().lower()))
        if mt is None:
            raise HTTPException(404, f"Machine type {machine_type} not found")
        return None, mt.id
    raise HTTPException(422, "Choose a machine or a machine type")


def upload(
    db: Session, user: User, file: UploadFile, title: str, kind: str, machine_no: str = "", machine_type: str = ""
) -> Document:
    if kind not in DOCUMENT_KINDS:
        raise HTTPException(422, f"kind must be one of {', '.join(DOCUMENT_KINDS)}")
    mime = (file.content_type or "").split(";")[0].strip().lower()
    if mime == "application/octet-stream" and (file.filename or "").lower().endswith((".csv", ".xlsx", ".md", ".txt")):
        mime = {".csv": "text/csv", ".md": "text/markdown", ".txt": "text/plain"}.get(
            Path(file.filename).suffix.lower(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    if mime not in MIMES:
        raise HTTPException(415, f"Unsupported file type: {mime or 'unknown'}")
    if kind == "alarm_list" and mime not in ("text/csv", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
        raise HTTPException(422, "An alarm list must be an Excel (.xlsx) or CSV file")
    machine_id, type_id = _scope(db, machine_no, machine_type)
    limit = get_config().max_document_mb * 1024 * 1024
    rel = Path("documents") / f"{uuid.uuid4().hex}{MIMES[mime]}"
    dest = get_config().media_dir / rel
    _dir()
    size = 0
    with dest.open("wb") as out:
        while chunk := file.file.read(1024 * 512):
            size += len(chunk)
            if size > limit:
                out.close()
                dest.unlink(missing_ok=True)
                raise HTTPException(413, f"File larger than {get_config().max_document_mb} MB")
            out.write(chunk)
    doc = Document(
        machine_id=machine_id,
        machine_type_id=type_id,
        title=(title or file.filename or "Document").strip()[:300],
        kind=kind,
        file_path=str(rel),
        mime=mime,
        size_bytes=size,
        uploaded_by=user.id,
    )
    db.add(doc)
    db.flush()
    audit.record(db, user, "document", doc.id, "create", None, audit.snapshot(doc))
    jobs.enqueue(db, "index_document", document_id=doc.id)
    return doc


def _chunks(text: str) -> list[str]:
    """Split on blank lines and pack paragraphs into chunks of about CHUNK_CHARS."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    out, current = [], ""
    for p in paragraphs:
        while len(p) > CHUNK_CHARS:  # a very long paragraph
            out.append(p[:CHUNK_CHARS])
            p = p[CHUNK_CHARS:]
        if len(current) + len(p) + 2 > CHUNK_CHARS and current:
            out.append(current)
            current = p
        else:
            current = f"{current}\n\n{p}" if current else p
    if current:
        out.append(current)
    return out


def _pages(doc: Document) -> list[str]:
    path = path_of(doc)
    if doc.mime == "application/pdf":
        from pypdf import PdfReader

        return [(page.extract_text() or "") for page in PdfReader(str(path)).pages]
    if doc.mime in ("text/plain", "text/markdown"):
        return [path.read_text(encoding="utf-8", errors="replace")]
    return []


def _alarm_rows(doc: Document) -> list[dict]:
    data = path_of(doc).read_bytes()
    if doc.mime == "text/csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    else:
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    headers = [re.sub(r"[\s_]+", " ", h.strip().lower()) for h in rows[0]]
    index = {}
    for key, aliases in ALARM_COLUMNS.items():
        for alias in aliases:
            if alias.replace("_", " ") in headers:
                index[key] = headers.index(alias.replace("_", " "))
                break
    if "code" not in index:
        raise ValueError("The first row needs an 'alarm_code' column")
    out = []
    for r in rows[1:]:
        values = {k: (r[i].strip() if i < len(r) and r[i] else "") for k, i in index.items()}
        if values.get("code") and not values.get("message", "").startswith("[SAMPLE] Delete"):
            out.append(values)
    return out


def _load_alarms(db: Session, doc: Document) -> int:
    count = 0
    machines = {m.machine_no.upper(): m for m in db.scalars(select(Machine))}
    types = {t.name.lower(): t for t in db.scalars(select(MachineType))}
    for row in _alarm_rows(doc):
        machine_id, type_id = doc.machine_id, doc.machine_type_id
        if row.get("machine_no") and row["machine_no"].upper() in machines:
            machine_id, type_id = machines[row["machine_no"].upper()].id, None
        elif row.get("machine_type") and row["machine_type"].lower() in types:
            machine_id, type_id = None, types[row["machine_type"].lower()].id
        same_scope = (AlarmCode.machine_id == machine_id) if machine_id else (
            (AlarmCode.machine_type_id == type_id) & AlarmCode.machine_id.is_(None)
        )
        for old in db.scalars(select(AlarmCode).where(AlarmCode.active.is_(True), func.upper(AlarmCode.code) == row["code"].upper(), same_scope)):
            old.active = False  # replaced by the newer list; kept for history
        db.add(AlarmCode(
            machine_id=machine_id, machine_type_id=type_id, code=row["code"], message=row.get("message", ""),
            cause=row.get("cause", ""), remedy=row.get("remedy", ""), document_id=doc.id,
        ))
        count += 1
    return count


def index(db: Session, doc: Document) -> None:
    """Extract text into chunks (and alarm codes for alarm lists). Safe to run again."""
    for old in list(doc.chunks):
        db.delete(old)  # derived data, rebuilt from the stored file
    db.flush()
    try:
        if doc.kind == "alarm_list":
            n = _load_alarms(db, doc)
            for seq, row in enumerate(_alarm_rows(doc)):
                text = f"Alarm {row['code']}: {row.get('message', '')}. Cause: {row.get('cause', '')}. Remedy: {row.get('remedy', '')}."
                db.add(DocumentChunk(document_id=doc.id, page=1, seq=seq, text=text))
            doc.pages = 1
            doc.status = "indexed" if n else "no_text"
        else:
            pages = _pages(doc)
            doc.pages = len(pages)
            seq = 0
            for page_no, text in enumerate(pages, start=1):
                for chunk in _chunks(text):
                    db.add(DocumentChunk(document_id=doc.id, page=page_no, seq=seq, text=chunk))
                    seq += 1
            doc.status = "indexed" if sum(len(p.strip()) for p in pages) >= MIN_TEXT else "no_text"
        doc.error = ""
    except Exception as e:  # a broken file must not stop the worker
        log.warning("Could not index document %s: %s", doc.id, e)
        doc.status = "failed"
        doc.error = str(e)[:500]
    doc.indexed_at = utcnow()


def set_active(db: Session, user: User, doc: Document, active: bool) -> Document:
    before = audit.snapshot(doc)
    doc.active = active
    if not active:
        for alarm in db.scalars(select(AlarmCode).where(AlarmCode.document_id == doc.id, AlarmCode.active.is_(True))):
            alarm.active = False
    audit.record(db, user, "document", doc.id, "update", before, audit.snapshot(doc))
    return doc


def visible_for(db: Session, machine: Machine) -> list[Document]:
    scope = Document.machine_id == machine.id
    if machine.machine_type_id:
        scope = scope | ((Document.machine_type_id == machine.machine_type_id) & Document.machine_id.is_(None))
    return list(db.scalars(select(Document).where(Document.active.is_(True), scope).order_by(Document.kind, Document.title)))


def serialize(db: Session, doc: Document) -> dict:
    machine = db.get(Machine, doc.machine_id) if doc.machine_id else None
    mtype = db.get(MachineType, doc.machine_type_id) if doc.machine_type_id else None
    return {
        "id": doc.id,
        "title": doc.title,
        "kind": doc.kind,
        "mime": doc.mime,
        "pages": doc.pages,
        "status": doc.status,
        "error": doc.error,
        "active": doc.active,
        "scope": "machine" if doc.machine_id else "type",
        "machine_no": machine.machine_no if machine else None,
        "type_name": mtype.name if mtype else "",
        "created_at": doc.created_at.isoformat(),
    }
