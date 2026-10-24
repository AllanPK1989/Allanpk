import io

from openpyxl import load_workbook

from .conftest import login


def test_dashboard_numbers_match_sample_data(client):
    sup = login(client, "D-SUP1")
    d = client.get("/api/dashboard?days=90", headers=sup).json()
    assert d["machines"] == 6
    # Sample data: 7 closed/verified logs, downtime 35+40+25+30+45+20+90.
    assert d["totals"]["breakdowns"] == 7
    assert d["totals"]["downtime_min"] == 285
    assert d["totals"]["mttr_min"] == round(285 / 7, 1)
    assert d["totals"]["logs_verified_pct"] == round(100 * 5 / 7)
    zones = {r["name"]: r for r in d["downtime_by_zone"]}
    assert zones["Bussmann Zone 1"]["downtime_min"] == 35 + 40 + 25 + 30 + 45 + 20
    assert zones["Zone 2"] == {"name": "Zone 2", "downtime_min": 90, "breakdowns": 1}
    shifts = {r["name"]: r["breakdowns"] for r in d["downtime_by_shift"]}
    assert sum(shifts.values()) == 7
    assert d["machines_table"][0]["machine_no"] == "DEMO-SF-01"  # most downtime first
    assert d["top_faults"][0]["symptom"] == "Cylinder not returning" and d["top_faults"][0]["count"] == 4
    k = d["knowledge"]
    assert k["machines_covered"] == 1 and len(k["machines_uncovered"]) == 5


def test_dashboard_zone_filter_and_roles(client):
    sup = login(client, "D-SUP1")
    d = client.get("/api/dashboard?days=90&zone=Zone 2", headers=sup).json()
    assert d["machines"] == 2 and d["totals"]["breakdowns"] == 1
    assert client.get("/api/dashboard", headers=login(client, "D-TECH1")).status_code == 403


def test_dashboard_excel_export(client):
    r = client.get("/api/export/dashboard.xlsx?days=90", headers=login(client, "D-ADMIN"))
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [
        "Summary", "Downtime by zone", "Downtime by shift", "Machines", "Top faults", "Spare parts", "PM overdue",
    ]
