import io
import json

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AICall, AuditLog, BreakdownLog

from .conftest import login


def test_login_lockout_after_five_wrong_pins(client):
    for _ in range(5):
        r = client.post("/api/auth/login", json={"employee_id": "D-TECH1", "pin": "0000"})
        assert r.status_code == 401
    r = client.post("/api/auth/login", json={"employee_id": "D-TECH1", "pin": "1234"})
    assert r.status_code == 423  # locked even with the right PIN


def test_roles_are_enforced(client):
    tech = login(client, "D-TECH1")
    assert client.get("/api/admin/users", headers=tech).status_code == 403
    assert client.get("/api/review/queue", headers=tech).status_code == 403
    assert client.get("/api/export/logs.xlsx", headers=tech).status_code == 403
    assert client.get("/api/machines").status_code == 401


def test_machine_page_shows_history_faults_and_kpis(client):
    h = login(client, "D-TECH1")
    page = client.get("/api/machines/demo-sf-01", headers=h).json()
    assert page["machine"]["machine_no"] == "DEMO-SF-01"
    assert page["kpi"]["breakdowns"] == 4 and page["kpi"]["mttr_min"] == 32.5
    assert page["top_faults"][0]["symptom"] == "Cylinder not returning"
    assert page["top_faults"][0]["count"] == 4
    assert [c["ref"] for c in page["fault_cards"]] == ["FC-1"]
    assert client.get("/api/machines/DEMO-SF-01/qr.svg", headers=h).text.startswith("<svg")


def test_ask_thunai_full_flow_creates_draft_log(client):
    h = login(client, "D-TECH1")
    s = client.post(
        "/api/ai/sessions",
        headers=h,
        json={"machine_no": "DEMO-SF-01", "problem_text": "Cylinder return aagala", "language": "ta"},
    ).json()
    assert s["outcome"] == "in_progress"
    assert 1 <= len(s["questions"]) <= 3
    answers = {q["id"]: q["options"][0]["value"] for q in s["questions"]}
    s = client.post(f"/api/ai/sessions/{s['id']}/answers", headers=h, json={"answers": answers}).json()
    steps = s["steps"]
    assert steps[0]["kind"] == "safety" and "LOTO" in steps[0]["action_en"]
    assert any("air" in st["action_en"].lower() for st in steps[:3])  # pneumatic machine
    cited = [st for st in steps if st["source_refs"]]
    assert cited and cited[0]["source_refs"] == ["FC-1"]
    assert any(st["authorised_only"] for st in steps)  # multimeter step
    r = client.post(f"/api/ai/sessions/{s['id']}/steps", headers=h, json={"step_no": 1, "response": "ok"})
    assert r.json()["next"] == 2
    r = client.post(f"/api/ai/sessions/{s['id']}/steps", headers=h, json={"step_no": 1, "response": "cant_do"})
    assert r.json()["next"] == -1  # cannot do a safety step → call supervisor
    sup = client.post(f"/api/ai/sessions/{s['id']}/call-supervisor", headers=h).json()
    assert "DEMO-SF-01" in sup["summary"] and sup["supervisors"]
    closed = client.post(f"/api/ai/sessions/{s['id']}/close", headers=h, json={"running": True}).json()
    log = client.get(f"/api/logs/{closed['log_id']}", headers=h).json()
    assert log["status"] == "draft" and log["needs_confirmation"]
    assert log["texts"]["issue"]["original"] == "Cylinder return aagala"  # technician's words kept
    assert log["texts"]["root_cause"]["ai_filled"]
    assert log["source"] == "ai_session"


def test_unsafe_request_is_refused_without_ai_call(client):
    h = login(client, "D-TECH1")
    s = client.post(
        "/api/ai/sessions",
        headers=h,
        json={"machine_no": "DEMO-NT-01", "problem_text": "How to bypass the light curtain interlock?", "language": "en"},
    ).json()
    assert s["outcome"] == "refused"
    assert s["steps"] == []
    assert s["procedure"]["call_supervisor"]
    with SessionLocal() as db:
        assert db.scalar(select(AICall).where(AICall.ai_session_id == s["id"])) is None


def test_escalation_keeps_log_open(client):
    h = login(client, "D-TECH2")
    s = client.post(
        "/api/ai/sessions", headers=h,
        json={"machine_no": "DEMO-AC-01", "problem_text": "Door seal leaking steam", "language": "en"},
    ).json()
    s = client.post(f"/api/ai/sessions/{s['id']}/answers", headers=h, json={"answers": {}}).json()
    assert any("pressure" in st["action_en"].lower() for st in s["steps"][:3])  # autoclave rule
    closed = client.post(f"/api/ai/sessions/{s['id']}/close", headers=h, json={"running": False}).json()
    assert closed["outcome"] == "escalated"
    log = client.get(f"/api/logs/{closed['log_id']}", headers=h).json()
    assert log["status"] == "open" and log["escalated"]


def test_text_log_submit_translate_verify_and_audit(client):
    tech = login(client, "D-TECH1")
    body = {
        "machine_no": "DEMO-SF-02",
        "issue": {"original": "Vibrator motor not running"},
        "root_cause": {"original": "MCB tripped"},
        "corrective_action": {"original": "Reset MCB, checked motor"},
        "downtime_min": 15,
        "parts": [{"description": "MCB 6A", "qty": 1}],
    }
    log = client.post("/api/logs", headers=tech, json=body).json()
    assert log["status"] == "closed" and log["shift"] in ("A", "B", "C")
    assert log["downtime_min"] == 15
    assert log["texts"]["issue"]["en"] == "Vibrator motor not running"
    # enrichment job ran after commit (sync mode) and classified the log
    log = client.get(f"/api/logs/{log['id']}", headers=tech).json()
    assert log["classification"] is not None
    sup = login(client, "D-SUP1")
    queue = client.get("/api/review/queue", headers=sup).json()
    assert any(l["id"] == log["id"] for l in queue["to_verify"])
    verified = client.post(f"/api/logs/{log['id']}/verify", headers=sup).json()
    assert verified["status"] == "verified"
    r = client.patch(f"/api/logs/{log['id']}", headers=tech, json={"alarm_code": "E1"})
    assert r.status_code == 403  # technicians cannot change verified logs
    with SessionLocal() as db:
        actions = [a.action for a in db.scalars(select(AuditLog).where(AuditLog.entity == "breakdown_log", AuditLog.entity_id == str(log["id"])))]
    assert actions == ["create", "verify"]


def test_submit_requires_issue_and_corrective_action(client):
    h = login(client, "D-TECH1")
    r = client.post("/api/logs", headers=h, json={"machine_no": "DEMO-SF-01", "issue": {"original": "Jam"}})
    assert r.status_code == 422
    assert "corrective_action" in r.json()["detail"]["missing"]


def test_voice_upload_and_extraction(client):
    h = login(client, "D-TECH1")
    r = client.post("/api/media/audio?language=ta", headers=h, files={"file": ("a.webm", b"fake-audio", "audio/webm")})
    assert r.status_code == 200 and r.json()["status"] == "transcribed"
    x = client.post("/api/logs/extract", headers=h, json={"machine_no": "DEMO-SF-01", "transcript": r.json()["transcript"]}).json()
    assert x["ok"] and x["corrective_action"]["original"]
    bad = client.post("/api/media/audio", headers=h, files={"file": ("a.exe", b"x", "application/octet-stream")})
    assert bad.status_code == 415


def test_offline_sync_is_idempotent_and_needs_confirmation(client):
    h = login(client, "D-TECH1")
    payload = json.dumps({"machine_no": "DEMO-SF-01", "client_uuid": "phone-1-log-7", "source": "voice"})
    files = {"audio": ("v.webm", b"audio", "audio/webm")}
    first = client.post("/api/logs/offline", headers=h, data={"payload": payload, "language": "ta"}, files=files).json()
    again = client.post("/api/logs/offline", headers=h, data={"payload": payload, "language": "ta"}, files=files).json()
    assert first["id"] == again["id"]
    log = client.get(f"/api/logs/{first['id']}", headers=h).json()
    assert log["needs_confirmation"] and log["texts"]["issue"]["ai_filled"]
    pending = client.get("/api/logs?scope=mine&needs_confirmation=true", headers=h).json()
    assert any(l["id"] == first["id"] for l in pending)


def test_excel_import_and_export(client):
    admin = login(client, "D-ADMIN")
    wb = Workbook()
    ws = wb.active
    ws.append(["Machine No", "Machine Name", "Zone", "Criticality"])
    ws.append(["DEMO-NEW-01", "[SAMPLE] New", "Zone 2", "c"])
    ws.append(["DEMO-BAD", "[SAMPLE] Bad", "Zone 2", "Z"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/admin/import/machines", headers=admin, files={"file": ("m.xlsx", buf.getvalue())}).json()
    assert r["created"] == 1 and len(r["errors"]) == 1

    wb = Workbook()
    ws = wb.active
    ws.append(["Machine No", "Date", "Start Time", "End Time", "Problem", "Action Taken"])
    ws.append(["DEMO-NEW-01", "12-10-2026", "22:30", "00:10", "Motor not running", "Reset MCB"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(
        "/api/admin/import/logs", headers=admin, files={"file": ("l.xlsx", buf.getvalue())}, data={"mark_verified": "true"}
    ).json()
    assert r["created"] == 1 and not r["errors"]
    with SessionLocal() as db:
        log = db.scalar(select(BreakdownLog).where(BreakdownLog.source == "import"))
        assert log.downtime_min == 100 and log.shift == "C" and log.status == "verified"

    x = client.get("/api/export/logs.xlsx", headers=admin)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Logs", "KPIs by machine"]
    assert wb["Logs"].max_row >= 9


def test_admin_user_management_and_settings(client):
    admin = login(client, "D-ADMIN")
    r = client.post(
        "/api/admin/users", headers=admin,
        json={"employee_id": "e100", "name": "New Tech", "role": "technician", "zone": "Zone 2", "pin": "4321"},
    )
    assert r.status_code == 200 and r.json()["employee_id"] == "E100"
    assert client.post("/api/auth/login", json={"employee_id": "E100", "pin": "4321"}).status_code == 200
    r = client.put("/api/admin/settings/shifts", headers=admin, json={"value": [{"name": "A", "start": "7:00"}]})
    assert r.status_code == 422
    r = client.put("/api/admin/settings/planned_hours_per_day", headers=admin, json={"value": 22.5})
    assert r.json()["planned_hours_per_day"] == 22.5


def test_api_has_no_delete_endpoints(client):
    spec = client.get("/api/openapi.json").json()
    methods = {m for ops in spec["paths"].values() for m in ops}
    assert "delete" not in methods
