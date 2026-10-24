import io

from openpyxl import Workbook, load_workbook

from app import parts
from app.models import SparePart

from .conftest import login


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_machine_page_shows_sample_parts_and_card_stock(client):
    page = client.get("/api/machines/DEMO-SF-01", headers=login(client, "D-TECH1")).json()
    stock = {p["part_no"]: p for p in page["spare_parts"]}
    assert set(stock) == {"DEMO-P-1001", "DEMO-P-1002", "DEMO-P-1003", "DEMO-P-1004"}  # sand filling type only
    assert stock["DEMO-P-1001"]["low"] and not stock["DEMO-P-1002"]["low"]  # 1 left, minimum 2
    assert stock["DEMO-P-1004"]["stock_qty"] == 0 and stock["DEMO-P-1004"]["low"]
    assert [p["part_no"] for p in page["fault_cards"][0]["parts_stock"]] == ["DEMO-P-1001"]
    lp = client.get("/api/machines/DEMO-LP-01", headers=login(client, "D-TECH1")).json()
    assert [p["part_no"] for p in lp["spare_parts"]] == ["DEMO-P-2001"]


def test_steps_show_stock_only_for_exact_part_numbers(client):
    tech = login(client, "D-TECH1")
    s = client.post("/api/ai/sessions", headers=tech, json={"machine_no": "DEMO-SF-01", "problem_text": "Cylinder not returning", "language": "en"}).json()
    if s["questions"]:
        s = client.post(f"/api/ai/sessions/{s['id']}/answers", headers=tech, json={"answers": {}}).json()
    with_parts = [st for st in s["steps"] if st["parts"]]
    assert len(with_parts) == 1 and "DEMO-P-1001" in with_parts[0]["action_en"]
    assert with_parts[0]["parts"][0]["location"] == "Store A-12" and with_parts[0]["parts"][0]["low"]


def test_find_in_text_is_exact_and_whole_word(db_session):
    part = db_session.query(SparePart).filter_by(part_no="DEMO-P-1001").one()
    part.alt_part_nos = ["MFR-COIL-24"]
    idx = parts.index([part])
    assert parts.find_in_text("Fit coil demo-p-1001 now.", idx) == [part]
    assert parts.find_in_text("Use MFR-COIL-24 (spare)", idx) == [part]
    assert parts.find_in_text("Part DEMO-P-10010 or DEMO-P-100", idx) == []  # no partial or fuzzy match


def test_parts_import_upserts_fits_and_reports_errors(client):
    admin = login(client, "D-ADMIN")
    content = xlsx([
        ["Material", "Material description", "Unrestricted", "Min stock", "Storage bin", "Machine type", "Mfr part no"],
        ["DEMO-P-1001", "[SAMPLE] Solenoid coil 24 V DC", 5, 2, "Store A-13", "Sand filling machine", "MFR-COIL-24"],
        ["DEMO-P-9001", "[SAMPLE] Label sensor", 3, 1, "Store C-02", "Label printer", ""],
        ["DEMO-P-9002", "[SAMPLE] Bad row", "lots", 1, "", "", ""],
        ["DEMO-P-9003", "[SAMPLE] Unknown type", 1, 1, "", "Rocket", ""],
        ["", "no part number", 1, 1, "", "", ""],
    ])
    r = client.post("/api/admin/import/parts", headers=admin, files={"file": ("stock.xlsx", content)})
    assert r.status_code == 200, r.text
    result = r.json()
    assert (result["created"], result["updated"]) == (1, 1)
    assert sorted(e["row"] for e in result["errors"]) == [4, 5, 6]
    page = client.get("/api/machines/DEMO-SF-01", headers=admin).json()
    coil = next(p for p in page["spare_parts"] if p["part_no"] == "DEMO-P-1001")
    assert coil["stock_qty"] == 5 and coil["location"] == "Store A-13" and not coil["low"]
    assert coil["alt_part_nos"] == ["MFR-COIL-24"]
    lp = client.get("/api/machines/DEMO-LP-01", headers=admin).json()
    assert {p["part_no"] for p in lp["spare_parts"]} == {"DEMO-P-2001", "DEMO-P-9001"}
    # CSV with semicolons (a common SAP export setting) works too.
    csv_bytes = "part_no;description;stock_qty\nDEMO-P-9001;[SAMPLE] Label sensor;0\n".encode()
    r = client.post("/api/admin/import/parts", headers=admin, files={"file": ("stock.csv", csv_bytes)})
    assert r.json()["updated"] == 1
    hits = client.get("/api/parts?q=label sensor&machine_no=DEMO-LP-01", headers=login(client, "D-TECH1")).json()
    assert hits[0]["part_no"] == "DEMO-P-9001" and hits[0]["stock_qty"] == 0
    assert client.get("/api/parts?q=%25", headers=admin).json() == []  # "%" is a character, not a wildcard
    t = client.get("/api/admin/templates/parts.xlsx", headers=admin)
    assert t.status_code == 200 and load_workbook(io.BytesIO(t.content)).active.max_row == 2  # header + SAMPLE row
    assert client.get("/api/admin/templates/pm.xlsx", headers=admin).status_code == 404


def test_dashboard_parts(client):
    sup = login(client, "D-SUP1")
    d = client.get("/api/dashboard?days=90", headers=sup).json()
    used = {r["description"]: r for r in d["parts"]["used"]}
    assert used["Solenoid coil"]["times"] == 3 and used["Solenoid coil"]["stock"] is None  # free-text history
    assert {p["part_no"] for p in d["parts"]["low_stock"]} == {"DEMO-P-1001", "DEMO-P-1004", "DEMO-P-3001"}
    assert "pm" not in d
