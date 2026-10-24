import io

from openpyxl import Workbook, load_workbook

from app import parts
from app.models import SparePart

from .conftest import login


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_machine_page_shows_sample_parts_pm_and_card_stock(client):
    page = client.get("/api/machines/DEMO-SF-01", headers=login(client, "D-TECH1")).json()
    stock = {p["part_no"]: p for p in page["spare_parts"]}
    assert set(stock) == {"DEMO-P-1001", "DEMO-P-1002", "DEMO-P-1003", "DEMO-P-1004"}  # sand filling type only
    assert stock["DEMO-P-1001"]["low"] and not stock["DEMO-P-1002"]["low"]  # 1 left, minimum 2
    assert stock["DEMO-P-1004"]["stock_qty"] == 0 and stock["DEMO-P-1004"]["low"]
    states = {t["task_code"]: (t["state"], t["days_overdue"]) for t in page["pm_tasks"]}
    assert states == {"PM-SF-01": ("overdue", 11), "PM-SF-02": ("due_soon", 0)}
    assert [p["part_no"] for p in page["fault_cards"][0]["parts_stock"]] == ["DEMO-P-1001"]
    lp = client.get("/api/machines/DEMO-LP-01", headers=login(client, "D-TECH1")).json()
    assert [p["part_no"] for p in lp["spare_parts"]] == ["DEMO-P-2001"]


def test_steps_show_stock_only_for_exact_part_numbers(client):
    tech = login(client, "D-TECH1")
    s = client.post("/api/ai/sessions", headers=tech, json={"machine_no": "DEMO-SF-01", "problem_text": "Cylinder not returning", "language": "en"}).json()
    if s["questions"]:
        s = client.post(f"/api/ai/sessions/{s['id']}/answers", headers=tech, json={"answers": {}}).json()
    with_parts = [st for st in s["steps"] if st["parts"]]
    assert len(with_parts) == 1 and "DEMO-P-1001" in with_parts[0]["action_en"]
    assert with_parts[0]["parts"][0]["location"] == "Store A-12" and with_parts[0]["parts"][0]["low"]
    # Overdue PM is evidence for the AI (tier 4) and cited by the step that uses it.
    pm_items = [x for x in s["sources"] if x["kind"] == "pm"]
    assert pm_items and pm_items[0]["ref"] == "PM-SF-01" and pm_items[0]["tier"] == 4
    assert any("PM-SF-01" in st["source_refs"] for st in s["steps"])


def test_find_in_text_is_exact_and_whole_word(db_session):
    part = db_session.query(SparePart).filter_by(part_no="DEMO-P-1001").one()
    part.alt_part_nos = ["MFR-COIL-24"]
    idx = parts.index([part])
    assert parts.find_in_text("Fit coil demo-p-1001 now.", idx) == [part]
    assert parts.find_in_text("Use MFR-COIL-24 (spare)", idx) == [part]
    assert parts.find_in_text("Part DEMO-P-10010 or DEMO-P-100", idx) == []  # no partial or fuzzy match


def test_parts_import_upserts_fits_and_reports_errors(client):
    admin = login(client, "D-ADMIN")
    content = xlsx([
        ["Material", "Material description", "Unrestricted", "Min stock", "Storage bin", "Machine type", "Mfr part no"],
        ["DEMO-P-1001", "[SAMPLE] Solenoid coil 24 V DC", 5, 2, "Store A-13", "Sand filling machine", "MFR-COIL-24"],
        ["DEMO-P-9001", "[SAMPLE] Label sensor", 3, 1, "Store C-02", "Label printer", ""],
        ["DEMO-P-9002", "[SAMPLE] Bad row", "lots", 1, "", "", ""],
        ["DEMO-P-9003", "[SAMPLE] Unknown type", 1, 1, "", "Rocket", ""],
        ["", "no part number", 1, 1, "", "", ""],
    ])
    r = client.post("/api/admin/import/parts", headers=admin, files={"file": ("stock.xlsx", content)})
    assert r.status_code == 200, r.text
    result = r.json()
    assert (result["created"], result["updated"]) == (1, 1)
    assert sorted(e["row"] for e in result["errors"]) == [4, 5, 6]
    page = client.get("/api/machines/DEMO-SF-01", headers=admin).json()
    coil = next(p for p in page["spare_parts"] if p["part_no"] == "DEMO-P-1001")
    assert coil["stock_qty"] == 5 and coil["location"] == "Store A-13" and not coil["low"]
    assert coil["alt_part_nos"] == ["MFR-COIL-24"]
    lp = client.get("/api/machines/DEMO-LP-01", headers=admin).json()
    assert {p["part_no"] for p in lp["spare_parts"]} == {"DEMO-P-2001", "DEMO-P-9001"}
    # CSV with semicolons (a common SAP export setting) works too.
    csv_bytes = "part_no;description;stock_qty\nDEMO-P-9001;[SAMPLE] Label sensor;0\n".encode()
    r = client.post("/api/admin/import/parts", headers=admin, files={"file": ("stock.csv", csv_bytes)})
    assert r.json()["updated"] == 1
    hits = client.get("/api/parts?q=label sensor&machine_no=DEMO-LP-01", headers=login(client, "D-TECH1")).json()
    assert hits[0]["part_no"] == "DEMO-P-9001" and hits[0]["stock_qty"] == 0
    assert client.get("/api/parts?q=%25", headers=admin).json() == []  # "%" is a character, not a wildcard


def test_pm_import_and_template(client):
    admin = login(client, "D-ADMIN")
    content = xlsx([
        ["Machine No", "Task code", "Task description", "Frequency (days)", "Last done", "Next due"],
        ["DEMO-LP-01", "PM-LP-01", "[SAMPLE] Clean print head", 14, "01-10-2026", ""],
        ["DEMO-LP-01", "PM-LP-02", "[SAMPLE] Check label sensor", 30, "", "2026-12-01"],
        ["DEMO-XX-99", "PM-1", "Unknown machine", 7, "", ""],
        ["DEMO-LP-01", "PM-LP-03", "Bad date", 7, "32-13-2026", ""],
    ])
    r = client.post("/api/admin/import/pm", headers=admin, files={"file": ("pm.xlsx", content)}).json()
    assert r["created"] == 2 and len(r["errors"]) == 2
    tasks = {t["task_code"]: t for t in client.get("/api/machines/DEMO-LP-01", headers=admin).json()["pm_tasks"]}
    assert tasks["PM-LP-01"]["next_due_on"] == "2026-10-15"  # last done + 14 days
    assert tasks["PM-LP-02"]["next_due_on"] == "2026-12-01"
    for kind in ("parts", "pm"):
        t = client.get(f"/api/admin/templates/{kind}.xlsx", headers=admin)
        assert t.status_code == 200
        assert load_workbook(io.BytesIO(t.content)).active.max_row == 2  # header + SAMPLE row


def test_dashboard_parts_and_pm(client):
    sup = login(client, "D-SUP1")
    d = client.get("/api/dashboard?days=90", headers=sup).json()
    used = {r["description"]: r for r in d["parts"]["used"]}
    assert used["Solenoid coil"]["times"] == 3 and used["Solenoid coil"]["stock"] is None  # free-text history
    assert {p["part_no"] for p in d["parts"]["low_stock"]} >= {"DEMO-P-1001", "DEMO-P-1004"}
    overdue = {(r["machine_no"], r["task_code"]): r["days_overdue"] for r in d["pm"]["overdue"]}
    assert overdue == {("DEMO-SF-01", "PM-SF-01"): 11, ("DEMO-AC-01", "PM-AC-01"): 5}
    assert d["pm"]["tasks"] == 4
