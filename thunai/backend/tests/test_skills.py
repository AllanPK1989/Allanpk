import io

from openpyxl import load_workbook

from app import settings_store

from .conftest import login


def test_skills_from_sample_data(client):
    d = client.get("/api/skills?days=90", headers=login(client, "D-SUP1")).json()
    assert d["zone"] == "Bussmann Zone 1"  # supervisors see their own zone
    people = {t["employee_id"]: t for t in d["technicians"]}
    t1, t3 = people["D-TECH1"], people["D-TECH3"]
    assert (t1["fixes"], t1["verified"]) == (6, 5) and t1["machine_types"] == ["Label printer", "Sand filling machine"]
    assert t1["median_downtime_min"] == 32 and t1["vs_team"] == 1.0
    assert (t3["fixes"], t3["sessions"], t3["escalated"], t3["escalation_pct"]) == (0, 3, 3, 100)
    assert t3["cant_do_authorised"] == 3
    assert d["matrix"] == [{"key": t1["key"], "name": t1["name"], "cells": {"Label printer": 1, "Sand filling machine": 5}}]
    types = {r["name"]: r for r in d["machine_types"]}
    assert types["Sand filling machine"]["fixers"] == 1 and types["Label printer"]["median_downtime_min"] is None  # 1 event < 3
    kinds = sorted((s["kind"], s["machine_type"]) for s in d["suggestions"])
    assert kinds == [("authorisation", None), ("pair", "Sand filling machine"),
                     ("single_person", "Label printer"), ("single_person", "Sand filling machine")]
    pair = next(s for s in d["suggestions"] if s["kind"] == "pair")
    assert pair["technician"] == t3["name"] and pair["partner"] == t1["name"]


def test_skills_minimum_events_and_access(client, db_session):
    settings_store.put(db_session, None, "skill_min_events", 10)
    db_session.commit()
    d = client.get("/api/skills?days=90", headers=login(client, "D-ADMIN")).json()
    assert d["zone"] == "" and d["min_events"] == 10
    assert all(t["median_downtime_min"] is None and t["escalation_pct"] is None for t in d["technicians"])
    assert not [s for s in d["suggestions"] if s["kind"] != "single_person"]
    zone2 = client.get("/api/skills?days=90&zone=Zone 2", headers=login(client, "D-ADMIN")).json()
    assert [t["employee_id"] for t in zone2["technicians"]] == ["D-TECH2"]
    # A supervisor cannot look at another zone's people.
    assert client.get("/api/skills?days=90&zone=Zone 2", headers=login(client, "D-SUP1")).json()["zone"] == "Bussmann Zone 1"
    assert client.get("/api/skills", headers=login(client, "D-TECH1")).status_code == 403


def test_skills_excel_export(client):
    r = client.get("/api/export/skills.xlsx?days=90", headers=login(client, "D-SUP1"))
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Technicians", "Coverage", "Machine types", "Suggestions", "About"]
    assert wb["Coverage"].max_row == 2 and wb["Suggestions"].max_row == 5
